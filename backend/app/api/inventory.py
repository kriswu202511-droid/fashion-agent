import io
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.dependencies import get_current_user
from app.database import get_db
from app.models.inventory import InventoryItem, Product
from app.models.user import User
from app.schemas.common import PaginatedResponse
from app.schemas.inventory import (
    InventoryItemCreate,
    InventoryItemResponse,
    ProductCreate,
    ProductResponse,
    ProductUpdate,
)
from app.services.cache import cache_service

router = APIRouter()


@router.get("/products", response_model=PaginatedResponse[ProductResponse])
async def list_products(
    page: int = 1,
    page_size: int = 20,
    category: Optional[str] = None,
    keyword: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    cache_key = cache_service.api_key(
        "inventory:products",
        current_user.id,
        page=page,
        page_size=page_size,
        category=category,
        keyword=keyword,
    )
    cached = await cache_service.get(cache_key)
    if cached is not None:
        return PaginatedResponse(**cached)

    query = select(Product).where(Product.user_id == current_user.id)
    count_query = select(func.count()).select_from(Product).where(Product.user_id == current_user.id)

    if category:
        query = query.where(Product.category == category)
        count_query = count_query.where(Product.category == category)
    if keyword:
        query = query.where(Product.name.contains(keyword) | Product.sku.contains(keyword))
        count_query = count_query.where(Product.name.contains(keyword) | Product.sku.contains(keyword))

    total = (await db.execute(count_query)).scalar() or 0
    query = query.order_by(Product.created_at.desc()).offset((page - 1) * page_size).limit(page_size)
    products = (await db.execute(query)).scalars().all()

    result = PaginatedResponse(
        items=[ProductResponse.model_validate(p) for p in products],
        total=total, page=page, page_size=page_size,
    )
    await cache_service.set(cache_key, result.model_dump(), ttl=60)
    return result


@router.post("/products", response_model=ProductResponse)
async def create_product(
    data: ProductCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    product = Product(user_id=current_user.id, **data.model_dump())
    db.add(product)
    await db.commit()
    await db.refresh(product)
    await cache_service.invalidate_prefix(f"api:inventory:products:{current_user.id}:")
    return product


@router.put("/products/{product_id}", response_model=ProductResponse)
async def update_product(
    product_id: str,
    data: ProductUpdate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    result = await db.execute(
        select(Product).where(Product.id == product_id, Product.user_id == current_user.id)
    )
    product = result.scalar_one_or_none()
    if not product:
        raise HTTPException(status_code=404, detail="商品不存在")

    for key, value in data.model_dump(exclude_unset=True).items():
        setattr(product, key, value)
    await db.commit()
    await db.refresh(product)
    await cache_service.invalidate_prefix(f"api:inventory:products:{current_user.id}:")
    return product


@router.delete("/products/{product_id}")
async def delete_product(
    product_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    result = await db.execute(
        select(Product).where(Product.id == product_id, Product.user_id == current_user.id)
    )
    product = result.scalar_one_or_none()
    if not product:
        raise HTTPException(status_code=404, detail="商品不存在")
    await db.delete(product)
    await db.commit()
    await cache_service.invalidate_prefix(f"api:inventory:products:{current_user.id}:")
    return {"message": "已删除"}


@router.post("/products/import-excel")
async def import_products_excel(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    import openpyxl

    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    headers = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]

    created = 0
    for row in rows:
        row_dict = dict(zip(headers, row))
        sku = str(row_dict.get("SKU", "") or "").strip()
        name = str(row_dict.get("名称", "") or row_dict.get("name", "") or "").strip()
        if not sku or not name:
            continue

        product = Product(
            user_id=current_user.id,
            sku=sku,
            name=name,
            category=str(row_dict.get("品类", "") or row_dict.get("category", "") or ""),
            price=float(row_dict.get("售价", 0) or row_dict.get("price", 0) or 0),
            cost=float(row_dict.get("成本", 0) or row_dict.get("cost", 0) or 0),
            supplier=str(row_dict.get("供应商", "") or row_dict.get("supplier", "") or ""),
            tags=str(row_dict.get("标签", "") or row_dict.get("tags", "") or ""),
        )
        db.add(product)
        created += 1

    await db.commit()
    return {"message": f"成功导入 {created} 个商品"}


@router.get("/products/export-excel")
async def export_products_excel(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    import openpyxl
    from io import BytesIO

    result = await db.execute(select(Product).where(Product.user_id == current_user.id))
    products = result.scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "商品列表"
    ws.append(["SKU", "名称", "品类", "子品类", "售价", "成本", "供应商", "标签", "状态"])

    for p in products:
        ws.append([p.sku, p.name, p.category, p.sub_category, p.price, p.cost, p.supplier, p.tags, p.status])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=products.xlsx"},
    )


@router.post("/inventory", response_model=InventoryItemResponse)
async def create_inventory_item(
    data: InventoryItemCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    item = InventoryItem(user_id=current_user.id, **data.model_dump())
    db.add(item)
    await db.commit()
    await db.refresh(item)
    return item


@router.get("/inventory", response_model=PaginatedResponse[InventoryItemResponse])
async def list_inventory(
    page: int = 1,
    page_size: int = 20,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    base_where = InventoryItem.user_id == current_user.id
    total = (await db.execute(
        select(func.count()).select_from(InventoryItem).where(base_where)
    )).scalar() or 0

    result = await db.execute(
        select(InventoryItem)
        .where(base_where)
        .order_by(InventoryItem.updated_at.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    )
    items = result.scalars().all()
    return PaginatedResponse(
        items=[InventoryItemResponse.model_validate(i) for i in items],
        total=total, page=page, page_size=page_size,
    )
